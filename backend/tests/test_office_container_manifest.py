"""readers/office.py's container-manifest detection.

Found on a real carrier's own public SI template, not our generator's
output (docs/EXTERNAL_VALIDATION.md): a container table with one row per
container and a column per attribute, which the ordinary "first cell is
the label, the rest is the value" table reader cannot read at all. Fixed
narrowly -- a table needs 3+ columns and a header row naming "container"
before this path is even reached, so it cannot change how any 2-column
label:value table already in the graded 520-email set is read.

Fixtures are built with `python-docx` in this file, never `data/bundle`
(CLAUDE.md rule 1 / docs/COLLABORATION.md), the same pattern
`backend/tests/test_api.py` uses for its own synthetic documents.
"""
from __future__ import annotations

import io

import docx
import pytest

from sdoc.readers.office import read_docx
from sdoc.schema import ParsedDoc


def _docx_bytes(build) -> bytes:
    d = docx.Document()
    build(d)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _add_table(d, rows: list[list[str]]) -> None:
    table = d.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    for r, values in enumerate(rows):
        for c, v in enumerate(values):
            table.rows[r].cells[c].text = v


def test_a_real_container_manifest_table_yields_a_count():
    """CMA CGM's own column names, a fictional shipment (docs/EXTERNAL_VALIDATION.md)."""
    data = _docx_bytes(lambda d: _add_table(d, [
        ["Nr", "Container Nr", "Seal Nr", "Number and Kind of packages",
         "Cargo Description", "Cargo Gross Weight IN KILOS"],
        ["1", "CMAU4471288", "SL9927341", "420 CARTONS", "ELECTRICAL COMPONENTS", "8,450"],
    ]))
    doc = read_docx(ParsedDoc(path="test.docx", ext=".docx"), data)
    counts = [c for c in doc.chunks if c.label == "Container Count"]
    assert len(counts) == 1, doc.chunks
    assert counts[0].value == "1"


def test_a_manifest_table_counts_every_data_row_not_just_the_first():
    data = _docx_bytes(lambda d: _add_table(d, [
        ["Nr", "Container Nr", "Seal Nr", "Packages", "Description", "Gross Weight"],
        ["1", "CMAU4471288", "SL9927341", "420 CARTONS", "ELECTRICAL COMPONENTS", "8,450"],
        ["2", "CMAU5512090", "SL9927342", "418 CARTONS", "ELECTRICAL COMPONENTS", "8,120"],
        ["3", "CMAU6603171", "SL9927343", "422 CARTONS", "ELECTRICAL COMPONENTS", "8,390"],
    ]))
    doc = read_docx(ParsedDoc(path="test.docx", ext=".docx"), data)
    counts = [c for c in doc.chunks if c.label == "Container Count"]
    assert len(counts) == 1, doc.chunks
    assert counts[0].value == "3"


def test_an_ordinary_two_column_table_is_completely_unaffected():
    """The shape every table in the graded 520-email set actually uses.

    Regression guard: this table has a "Container No." row, same word the
    detector keys on, but only two columns -- must still resolve exactly as
    it always has, not be swept into manifest handling.
    """
    data = _docx_bytes(lambda d: _add_table(d, [
        ["Shipper", "TEST EXPORT COMPANY LTD"],
        ["Container No.", "TCLU1234567"],
        ["Total Containers", "2 x 40'HC"],
    ]))
    doc = read_docx(ParsedDoc(path="test.docx", ext=".docx"), data)
    assert not any(c.label == "Container Count" for c in doc.chunks)
    by_label = {c.label: c.value for c in doc.chunks}
    assert by_label["Shipper"] == "TEST EXPORT COMPANY LTD"
    assert by_label["Container No."] == "TCLU1234567"
    assert by_label["Total Containers"] == "2 x 40'HC"


def test_a_wide_table_that_is_not_about_containers_is_not_a_manifest():
    """3+ columns alone is not the signal -- the header must actually name one."""
    data = _docx_bytes(lambda d: _add_table(d, [
        ["Item", "Unit Price", "Quantity", "Amount"],
        ["Handling fee", "50.00", "1", "50.00"],
    ]))
    doc = read_docx(ParsedDoc(path="test.docx", ext=".docx"), data)
    assert not any(c.label == "Container Count" for c in doc.chunks)


@pytest.mark.parametrize("header_word", ["Container Nr", "CONTAINER NO", "container number"])
def test_the_detector_is_case_and_punctuation_insensitive(header_word):
    data = _docx_bytes(lambda d: _add_table(d, [
        [header_word, "Seal Nr", "Description"],
        ["CMAU4471288", "SL9927341", "ELECTRICAL COMPONENTS"],
    ]))
    doc = read_docx(ParsedDoc(path="test.docx", ext=".docx"), data)
    assert any(c.label == "Container Count" for c in doc.chunks)
