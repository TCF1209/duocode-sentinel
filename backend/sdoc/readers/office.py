"""Word (.docx) and Excel (.xlsx) attachment readers.

Both formats already carry the label/value structure we want:

    xlsx   column A = label, column B = value
    docx   a 2-column "Table Grid", one field per row

so reading them is mostly about producing a useful locator ("cell A7",
"table 1 row 3") and flattening the value to a string.
"""
from __future__ import annotations

import io

from .. import normalize
from ..schema import Chunk, ParsedDoc
from .rows import split_label_value

# A real container-manifest table -- one row per container, a column per
# attribute (Nr / Container Nr / Seal Nr / packages / description / weight) --
# is not the "label in column 1, value in the rest" shape the loop below
# assumes; docs/EXTERNAL_VALIDATION.md found one on a real carrier's own SI
# template, where it meant `container_count` came back missing rather than
# guessed (safe, but not useful). Detected narrowly, on the header row alone,
# so the ordinary 2-column label:value table -- every table in the graded
# 520-email set -- is completely unaffected; that shape never has 3+ columns.
_MIN_MANIFEST_COLUMNS = 3


# --------------------------------------------------------------------------
# .xlsx
# --------------------------------------------------------------------------
def read_xlsx(doc: ParsedDoc, data: bytes) -> ParsedDoc:
    import openpyxl

    if not data:
        doc.readable = False
        doc.unreadable_reason = "empty_file"
        return doc

    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    except Exception as exc:                      # corrupt / not really xlsx
        doc.readable = False
        doc.unreadable_reason = "corrupt"
        doc.notes.append(f"openpyxl: {type(exc).__name__}: {exc}")
        return doc

    lines: list[str] = []
    chunks: list[Chunk] = []
    order = 0
    for ws in wb.worksheets:
        for r, row in enumerate(ws.iter_rows(values_only=True), start=1):
            cells = ["" if c is None else str(c).strip() for c in row]
            while cells and not cells[-1]:
                cells.pop()
            if not cells:
                continue
            lines.append("\t".join(cells))

            label = cells[0]
            value = " ".join(cells[1:]).strip()
            if label and value:
                chunks.append(
                    Chunk(label=label, value=value,
                          locator=f"{ws.title}!A{r}", order=order)
                )
                order += 1
            elif label:
                # a single-cell row can still be "Label: value" written inline
                pair = split_label_value(label)
                if pair:
                    chunks.append(
                        Chunk(label=pair[0], value=pair[1],
                              locator=f"{ws.title}!A{r}", order=order)
                    )
                    order += 1

    wb.close()
    doc.text = "\n".join(lines)
    doc.chunks = chunks
    if not doc.text.strip():
        doc.readable = False
        doc.unreadable_reason = "empty_file"
    return doc


def _is_container_manifest_header(header_cells: list[str]) -> bool:
    """Does this table's first row name a container column -- "Container Nr", "Container No", etc.

    Only reached for a table with 3+ columns already (`_MIN_MANIFEST_COLUMNS`),
    so this never looks at the 2-column label:value tables every document in
    the graded 520-email set actually uses. Deliberately just the one word:
    "CONTAINER" is what CMA CGM's own public SI template calls this column,
    and what any other carrier's manifest would too, in every spelling
    `normalize.basic` already folds together (case, punctuation).
    """
    return any("CONTAINER" in normalize.basic(cell) for cell in header_cells)


# --------------------------------------------------------------------------
# .docx
# --------------------------------------------------------------------------
def read_docx(doc: ParsedDoc, data: bytes) -> ParsedDoc:
    import docx

    if not data:
        doc.readable = False
        doc.unreadable_reason = "empty_file"
        return doc

    try:
        d = docx.Document(io.BytesIO(data))
    except Exception as exc:
        doc.readable = False
        doc.unreadable_reason = "corrupt"
        doc.notes.append(f"python-docx: {type(exc).__name__}: {exc}")
        return doc

    lines: list[str] = []
    chunks: list[Chunk] = []
    order = 0

    for p in d.paragraphs:
        t = p.text.strip()
        if not t:
            continue
        lines.append(t)
        pair = split_label_value(t)
        if pair:
            chunks.append(Chunk(label=pair[0], value=pair[1],
                                locator=f"para {len(lines)}", order=order))
            order += 1

    for ti, table in enumerate(d.tables, start=1):
        is_manifest = False    # this table is one row per container, not label:value
        container_rows = 0
        for ri, row in enumerate(table.rows, start=1):
            cells = [c.text.strip() for c in row.cells]
            # python-docx repeats merged cells; collapse consecutive duplicates
            dedup: list[str] = []
            for c in cells:
                if not dedup or dedup[-1] != c:
                    dedup.append(c)
            if not any(dedup):
                continue
            lines.append(" | ".join(dedup))

            if ri == 1 and len(dedup) >= _MIN_MANIFEST_COLUMNS and _is_container_manifest_header(dedup):
                # A header row naming containers, not a 2-column label:value
                # row -- every row after this one is one container, counted
                # below rather than misread as its own label/value pair.
                is_manifest = True
                continue
            if is_manifest:
                container_rows += 1
            elif len(dedup) >= 2 and dedup[0]:
                chunks.append(
                    Chunk(label=dedup[0], value=" ".join(x for x in dedup[1:] if x).strip(),
                          locator=f"table {ti} row {ri}", order=order)
                )
                order += 1

        if is_manifest and container_rows > 0:
            chunks.append(
                Chunk(label="Container Count", value=str(container_rows),
                      locator=f"table {ti} ({container_rows} rows)", order=order)
            )
            order += 1

    doc.text = "\n".join(lines)
    doc.chunks = chunks
    if not doc.text.strip():
        doc.readable = False
        doc.unreadable_reason = "empty_file"
    return doc
